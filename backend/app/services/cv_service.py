import re
from io import BytesIO
from typing import List, Optional

from starlette.concurrency import run_in_threadpool

from app.core.errors import ServiceError
from app.data.skill_taxonomy import category_of, extract_skills
from app.schemas.cv import CandidateSkill, CVExtractionResponse

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Số điện thoại Việt Nam: 0xxxxxxxxx hoặc +84xxxxxxxxx, cho phép dấu cách . -
_PHONE = re.compile(r"(?<!\d)(?:\+?84|0)[\s.-]?\d{2,3}[\s.-]?\d{3}[\s.-]?\d{3,4}(?!\d)")


class CVService:
    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(file_bytes))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as exc:
            raise ServiceError(422, f"Không đọc được file PDF: {exc}") from exc

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        from docx import Document

        try:
            document = Document(BytesIO(file_bytes))
        except Exception as exc:
            raise ServiceError(422, f"Không đọc được file DOCX: {exc}") from exc
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return "\n".join(lines)

    @staticmethod
    def extract_text(file_bytes: bytes, filename: str) -> str:
        name = filename.lower()
        if name.endswith(".pdf"):
            text = CVService._extract_pdf(file_bytes)
        elif name.endswith(".docx"):
            text = CVService._extract_docx(file_bytes)
        else:
            raise ServiceError(400, "Chỉ hỗ trợ file định dạng .pdf hoặc .docx")
        if not text.strip():
            raise ServiceError(
                422, "File không chứa văn bản đọc được (có thể là CV dạng ảnh scan, cần OCR)."
            )
        return text

    @staticmethod
    def _guess_name(text: str) -> Optional[str]:
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            # Tên thường là dòng đầu tiên: ngắn, không số, không email.
            if len(line.split()) <= 5 and not re.search(r"[\d@]", line):
                return line
            break
        return None

    @staticmethod
    def parse_text(text: str, filename: str) -> CVExtractionResponse:
        email = _EMAIL.search(text)
        phone = _PHONE.search(text)
        skills = [CandidateSkill(name=s, category=category_of(s)) for s in extract_skills(text)]

        warnings: List[str] = [
            "Học vấn và kinh nghiệm chưa được bóc tách (cần tích hợp LLM/NER tại cv_service.parse_text)."
        ]
        if not skills:
            warnings.append("Không nhận diện được kỹ năng nào trong Skill Taxonomy.")

        return CVExtractionResponse(
            full_name=CVService._guess_name(text),
            email=email.group(0) if email else None,
            phone=re.sub(r"[\s.-]", "", phone.group(0)) if phone else None,
            skills=skills,
            raw_summary=f"Đã đọc {len(text)} ký tự từ file {filename}.",
            is_mock=True,
            warnings=warnings,
        )

    @staticmethod
    async def parse_cv(file_bytes: bytes, filename: str) -> CVExtractionResponse:
        """
        Đọc file (PDF/DOCX) và bóc tách thông tin.
        Hiện: email, số điện thoại, họ tên (heuristic) và kỹ năng (theo Skill Taxonomy) là kết quả thật.
        TODO (nhóm AI): dùng prompt `cv_extraction` + gemini_service.generate_json để bóc tách học vấn,
        kinh nghiệm, rồi bỏ cờ is_mock.
        """
        text = await run_in_threadpool(CVService.extract_text, file_bytes, filename)
        return CVService.parse_text(text, filename)


cv_service = CVService()
